"""
Reading and writing .csv files for the Document Redactor -- shared by the
spreadsheet redaction plugin (document_types/excel_spreadsheet.py) and the
spreadsheet un-redact engine (spreadsheet_unredact_engine.py).

WHY A SHARED MODULE
A redacted (or un-redacted) .csv should differ from its original ONLY in
the cells that were changed. Practice-management and bank exports come in
several dialects, so this module notes how a file was written and writes
the result back the same way:
  * encoding  -- UTF-8 with a byte-order mark (what Excel's "CSV UTF-8"
                 writes), plain UTF-8, UTF-16 (some exports), or Windows
                 code page 1252 (older Windows programs; accented names).
  * delimiter -- comma, semicolon, tab or pipe, chosen as the one that
                 splits the first lines into the same number of fields.
  * quoting   -- quotes only where needed (normal), or every field quoted
                 when the original quoted every field.
  * line ends -- Windows (CRLF) or Unix (LF), and whether the last line
                 ended with one.
Every value is kept as TEXT, exactly as written: "00123" stays "00123",
"1.50" stays "1.50", a date stays in its own format. Nothing is converted
to a number or date the way opening a .csv in Excel would.

AS A WORKBOOK
load_as_workbook() presents a .csv to the spreadsheet plugin as a one-sheet
openpyxl workbook (sheet name "CSV"), so the same column-picking,
combining, numbering, key and verification code handles .xlsx and .csv
alike. save_workbook() writes such a workbook back as .csv in the
original's dialect (a normal workbook is saved as .xlsx as usual).
"""
import csv
import io
import os
from dataclasses import dataclass, field
from typing import List

SHEET_NAME = "CSV"
_DELIMITERS = [",", "\t", ";", "|"]


@dataclass
class CsvFormat:
    encoding: str = "utf-8"
    delimiter: str = ","
    quote_all: bool = False
    lineterminator: str = "\r\n"
    trailing_newline: bool = True
    row_lengths: List[int] = field(default_factory=list)


def is_csv(path) -> bool:
    return os.path.splitext(str(path))[1].lower() in (".csv", ".tsv")


def _decode(raw: bytes):
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8"), "utf-8-sig"
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), "utf-16"
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace"), "cp1252"


def _pick_delimiter(text: str, path: str) -> str:
    sample_lines = text[:65536].splitlines()[:50]
    sample = "\n".join(sample_lines)
    best, best_score = None, 0.0
    for d in _DELIMITERS:
        try:
            counts = [len(r) for r in csv.reader(io.StringIO(sample), delimiter=d) if r]
        except csv.Error:
            continue
        if not counts or counts[0] < 2:
            continue
        score = sum(1 for c in counts if c == counts[0]) / len(counts)
        if score > best_score:
            best, best_score = d, score
    if best:
        return best
    return "\t" if path.lower().endswith(".tsv") else ","


def _all_quoted(line: str, delimiter: str) -> bool:
    line = line.rstrip("\r\n")
    if len(line) < 2 or not (line.startswith('"') and line.endswith('"')):
        return False
    try:
        fields = next(csv.reader([line], delimiter=delimiter))
    except (csv.Error, StopIteration):
        return False
    return line.count('"' + delimiter + '"') >= len(fields) - 1


def read_csv(path: str):
    """(rows, CsvFormat). rows is a list of lists of str."""
    try:
        with open(path, "rb") as f:
            raw = f.read()
    except OSError as exc:
        raise RuntimeError(f"Could not read this CSV file: {exc}") from exc
    text, encoding = _decode(raw)
    delimiter = _pick_delimiter(text, path)
    first_lines = text[:65536].splitlines()[:3]
    fmt = CsvFormat(
        encoding=encoding,
        delimiter=delimiter,
        quote_all=bool(first_lines) and all(_all_quoted(l, delimiter) for l in first_lines if l),
        lineterminator="\r\n" if "\r\n" in text[:65536] or (not "\n" in text[:65536]) else "\n",
        trailing_newline=text.endswith(("\n", "\r")) or not text,
    )
    try:
        rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    except csv.Error as exc:
        raise RuntimeError(f"Could not read this CSV file ({exc}). Open it in Excel and save a fresh copy as CSV.") from exc
    fmt.row_lengths = [len(r) for r in rows]
    return rows, fmt


def write_csv(path: str, rows, fmt: CsvFormat = None) -> None:
    fmt = fmt or CsvFormat()
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, delimiter=fmt.delimiter, lineterminator=fmt.lineterminator,
                        quoting=csv.QUOTE_ALL if fmt.quote_all else csv.QUOTE_MINIMAL)
    for r in rows:
        writer.writerow(["" if v is None else v for v in r])
    text = buf.getvalue()
    if not fmt.trailing_newline and text.endswith(fmt.lineterminator):
        text = text[: -len(fmt.lineterminator)]
    try:
        data = text.encode(fmt.encoding)
        if fmt.encoding == "utf-8-sig" and not data.startswith(b"\xef\xbb\xbf"):
            data = b"\xef\xbb\xbf" + data
    except UnicodeEncodeError:
        # A restored value the original code page can't hold -- UTF-8 with
        # a byte-order mark keeps every character and Excel reads it.
        data = b"\xef\xbb\xbf" + text.encode("utf-8")
    with open(path, "wb") as f:
        f.write(data)


def load_as_workbook(path: str):
    """A one-sheet openpyxl Workbook holding the CSV's cells as text (blank
    cells as None). The file's dialect rides along as wb._csv_format."""
    import openpyxl
    from openpyxl.cell.cell import TYPE_STRING

    rows, fmt = read_csv(path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for r_i, row in enumerate(rows, start=1):
        for c_i, v in enumerate(row, start=1):
            if v == "":
                continue
            cell = ws.cell(row=r_i, column=c_i)
            try:
                cell.value = v
            except Exception:  # characters openpyxl refuses (control codes)
                cell._value = v
            cell.data_type = TYPE_STRING  # never a formula, number or date
    wb._csv_format = fmt
    return wb


def save_workbook(wb, path: str) -> None:
    """Saves `wb` to `path`: as .csv in the dialect it was loaded with when
    `path` is a .csv, otherwise as a normal .xlsx workbook."""
    if not is_csv(path):
        wb.save(path)
        return
    fmt = getattr(wb, "_csv_format", None) or CsvFormat()
    ws = wb.worksheets[0]
    max_c = ws.max_column if ws._cells else 0
    n_rows = max(len(fmt.row_lengths), ws.max_row if ws._cells else 0)
    out = []
    for r_i in range(1, n_rows + 1):
        orig_len = fmt.row_lengths[r_i - 1] if r_i - 1 < len(fmt.row_lengths) else 0
        values = []
        last = 0
        for c_i in range(1, max(orig_len, max_c) + 1):
            cell = ws._cells.get((r_i, c_i))
            v = None if cell is None else cell.value
            values.append("" if v is None else str(v))
            if v not in (None, ""):
                last = c_i
        out.append(values[: max(orig_len, last)])
    write_csv(path, out, fmt)
