"""
PDF to Excel -- the conversion engine (no user interface; see
pdf_to_excel_app.py for the window).

Turns the text of a PDF -- a bank statement, a practice-management report,
any page laid out in columns -- into rows and columns for Excel, using
column breaks the user places by hand (or that auto_detect_breaks()
suggests). Works on PDFs that contain real text. A scanned page is only a
picture of text and has none; page_has_text() says so and the window
tells the user (no OCR here).

HOW A PAGE BECOMES ROWS
  1. Words come from PyMuPDF with their positions (page_words()), already
     turned the way the page is displayed, so they line up with what the
     user sees on screen even on a rotated page.
  2. Only words inside the TABLE AREA count -- between the top and bottom
     bars (and, optionally, left/right), so page headers, logos, running
     totals and footers can be left out.
  3. Words are grouped into LINES by their vertical position: a word joins
     a line when its middle sits within half a typical text height of the
     line's middle, which keeps a slightly raised or lowered word (a
     superscript, a mixed font) on its own line rather than starting a new
     row.
  4. Each word goes into the COLUMN its horizontal middle falls in, between
     the column breaks. Words landing in the same cell are joined with a
     space, left to right.

OPTIONS (Options dataclass)
  join_wrapped   -- a line whose first column is empty is treated as the
                    continuation of the row above (a description that wraps
                    onto a second line in a bank statement); its cells are
                    appended to that row's cells. Off when every row should
                    stay its own row.
  skip_repeated_header -- after the first page, a row identical to the
                    first row of the whole extraction (the column headings
                    reprinted on every page) is dropped.
  convert_numbers -- "1,234.56", "$1,234.56", "(45.00)", "45.00-", "-45"
                    become numbers (negative where shown so). Anything that
                    could be an identifier is left as text: leading zeros
                    ("00123"), more than 15 digits, or no digits at all.
                    Dates are always left as text, so Excel never guesses a
                    day/month order.

LAYOUTS
  A Layout (column breaks + table area + options) can be saved as a small
  JSON file and loaded again for the next statement in the same format.
  Positions are stored in PDF points from the page's top-left corner, so a
  layout is independent of screen size and zoom.
"""
import json
import re
import statistics
from dataclasses import dataclass, field, asdict
from typing import Dict, List, Optional

try:
    import pymupdf as fitz  # PyMuPDF
except ImportError:  # older PyMuPDF
    import fitz

LAYOUT_VERSION = 1


@dataclass
class Word:
    x0: float
    y0: float
    x1: float
    y1: float
    text: str

    @property
    def xc(self):
        return (self.x0 + self.x1) / 2

    @property
    def yc(self):
        return (self.y0 + self.y1) / 2

    @property
    def height(self):
        return self.y1 - self.y0


@dataclass
class Area:
    """The table area on a page, in points. None = page edge."""
    top: Optional[float] = None
    bottom: Optional[float] = None
    left: Optional[float] = None
    right: Optional[float] = None

    def contains(self, w: Word) -> bool:
        if self.top is not None and w.yc < self.top:
            return False
        if self.bottom is not None and w.yc > self.bottom:
            return False
        if self.left is not None and w.xc < self.left:
            return False
        if self.right is not None and w.xc > self.right:
            return False
        return True


@dataclass
class Options:
    join_wrapped: bool = True
    skip_repeated_header: bool = True
    convert_numbers: bool = True


@dataclass
class Layout:
    breaks: List[float] = field(default_factory=list)      # x positions, points
    area: Area = field(default_factory=Area)                # same on every page...
    page_areas: Dict[int, Area] = field(default_factory=dict)  # ...unless overridden here (0-based page)
    options: Options = field(default_factory=Options)

    def area_for(self, page_index: int) -> Area:
        return self.page_areas.get(page_index, self.area)

    # -- saving / loading ---------------------------------------------
    def to_json(self) -> str:
        data = {
            "layout_version": LAYOUT_VERSION,
            "breaks": sorted(self.breaks),
            "area": asdict(self.area),
            "page_areas": {str(k): asdict(v) for k, v in self.page_areas.items()},
            "options": asdict(self.options),
        }
        return json.dumps(data, indent=2)

    @classmethod
    def from_json(cls, text: str) -> "Layout":
        data = json.loads(text)
        if not isinstance(data, dict) or "breaks" not in data:
            raise ValueError("This is not a PDF to Excel layout file.")
        opts = Options(**{k: v for k, v in (data.get("options") or {}).items() if k in Options.__dataclass_fields__})
        return cls(
            breaks=sorted(float(b) for b in data.get("breaks", [])),
            area=Area(**(data.get("area") or {})),
            page_areas={int(k): Area(**v) for k, v in (data.get("page_areas") or {}).items()},
            options=opts,
        )


# --------------------------------------------------------------------------
# Reading the PDF
# --------------------------------------------------------------------------
def open_pdf(path: str):
    doc = fitz.open(path)
    if doc.needs_pass:
        doc.close()
        raise ValueError("This PDF is password-protected. Open it in a PDF viewer, save an unprotected copy, and use that.")
    return doc


def page_words(page) -> List[Word]:
    """Every word on the page with its box, in DISPLAYED orientation
    (rotation applied), top-left origin, in points."""
    rot = page.rotation_matrix
    out = []
    for x0, y0, x1, y1, text, *_rest in page.get_text("words"):
        if not text.strip():
            continue
        r = fitz.Rect(x0, y0, x1, y1) * rot
        out.append(Word(r.x0, r.y0, r.x1, r.y1, text))
    return out


def page_has_text(page) -> bool:
    return bool(page.get_text("words"))


def page_size(page):
    """(width, height) as displayed (page.rect already has the page's
    rotation applied; the words need page.rotation_matrix, see above)."""
    return page.rect.width, page.rect.height


# --------------------------------------------------------------------------
# Rows and columns
# --------------------------------------------------------------------------
def group_lines(words: List[Word]) -> List[List[Word]]:
    """Words -> lines (top to bottom), each line left to right."""
    if not words:
        return []
    heights = [w.height for w in words if w.height > 0]
    tol = (statistics.median(heights) if heights else 8.0) * 0.5
    lines = []
    for w in sorted(words, key=lambda w: (w.yc, w.x0)):
        if lines and abs(w.yc - lines[-1]["yc"]) <= tol:
            line = lines[-1]
            line["words"].append(w)
            line["yc"] = sum(x.yc for x in line["words"]) / len(line["words"])
        else:
            lines.append({"yc": w.yc, "words": [w]})
    return [sorted(l["words"], key=lambda w: w.x0) for l in lines]


def split_line(line: List[Word], breaks: List[float]) -> List[str]:
    """One line -> cells, by the column breaks."""
    bs = sorted(breaks)
    cells = [[] for _ in range(len(bs) + 1)]
    for w in line:
        col = sum(1 for b in bs if w.xc > b)
        cells[col].append(w.text)
    return [" ".join(c) for c in cells]


def page_rows(page, layout: Layout, page_index: int) -> List[List[str]]:
    """The rows of one page with the layout applied (before cross-page
    options such as skipping repeated headers)."""
    area = layout.area_for(page_index)
    words = [w for w in page_words(page) if area.contains(w)]
    rows = [split_line(line, layout.breaks) for line in group_lines(words)]
    if layout.options.join_wrapped:
        rows = join_wrapped_rows(rows)
    return rows


def join_wrapped_rows(rows: List[List[str]]) -> List[List[str]]:
    out = []
    for row in rows:
        if out and not row[0].strip() and any(c.strip() for c in row[1:]):
            prev = out[-1]
            for i, cell in enumerate(row):
                if cell.strip():
                    prev[i] = (prev[i] + " " + cell).strip() if prev[i] else cell
        else:
            out.append(list(row))
    return out


def extract(doc, layout: Layout, pages: List[int], progress=None) -> Dict[int, List[List[str]]]:
    """{page_index: rows} for the given pages, with every option applied."""
    result = {}
    header = None
    for n, i in enumerate(pages, 1):
        rows = page_rows(doc[i], layout, i)
        if layout.options.skip_repeated_header:
            if header is None and rows:
                header = [c.strip() for c in rows[0]]
            elif header is not None:
                rows = [r for r in rows if [c.strip() for c in r] != header]
        result[i] = rows
        if progress:
            progress(n, len(pages))
    return result


# --------------------------------------------------------------------------
# Suggesting column breaks
# --------------------------------------------------------------------------
def auto_detect_breaks(words: List[Word], area: Area = None, min_gap: float = 7.0,
                       min_rows: int = 3, tolerance: float = 0.2) -> List[float]:
    """Suggests column breaks from vertical white space: strips at least
    `min_gap` points wide, with text on both sides, that almost no line
    crosses. "Almost" (up to `tolerance` of the lines, e.g. 3 lines in 15)
    lets a few lines that run across the columns -- a bank's name and
    address, a heading, a long memo -- sit above or inside the table
    without hiding every gap. Lines with one word (a page number) are
    ignored. A suggestion, not a guarantee -- the user drags, adds or
    removes breaks afterwards."""
    area = area or Area()
    words = [w for w in words if area.contains(w)]
    lines = [l for l in group_lines(words) if len(l) >= 2]
    if len(lines) < min_rows:
        lines = group_lines(words)
    used = [w for l in lines for w in l]
    if not used:
        return []
    lo = int(min(w.x0 for w in used))
    hi = int(max(w.x1 for w in used)) + 1
    covered = [0] * (hi - lo + 1)
    for line in lines:
        hit = set()
        for w in line:
            hit.update(range(int(w.x0) - lo, int(w.x1) - lo + 1))
        for x in hit:
            covered[x] += 1
    allowed = int(len(lines) * tolerance) if len(lines) >= 6 else 0
    breaks = []
    x = 0
    while x < len(covered):
        if covered[x] <= allowed:
            start = x
            while x < len(covered) and covered[x] <= allowed:
                x += 1
            if x - start >= min_gap and start > 0 and x < len(covered):
                # center on the emptiest part of the strip
                seg = covered[start:x]
                best = min(seg)
                idx = [start + k for k, v in enumerate(seg) if v == best]
                breaks.append(lo + (idx[0] + idx[-1]) / 2 + 0.5)
        else:
            x += 1
    return breaks


# --------------------------------------------------------------------------
# Numbers
# --------------------------------------------------------------------------
_NUM_RE = re.compile(r"^(?P<neg1>\()?\s*(?P<sign>[-+])?\s*(?P<cur>[$€£])?\s*(?P<num>\d{1,3}(?:,\d{3})+|\d+)(?P<dec>\.\d+)?\s*(?P<neg2>\))?(?P<trail>-|CR|DR)?$", re.I)


def to_number(text: str):
    """The number shown by `text`, or None to keep it as text (see module
    docstring for what is deliberately left alone)."""
    s = text.strip()
    if not s:
        return None
    m = _NUM_RE.match(s)
    if not m:
        return None
    if bool(m.group("neg1")) != bool(m.group("neg2")):
        return None
    digits = m.group("num").replace(",", "")
    if len(digits) > 15:
        return None
    if len(digits) > 1 and digits.startswith("0") and not m.group("dec"):
        return None  # account / check numbers with leading zeros stay text
    value = float(digits + (m.group("dec") or ""))
    negative = bool(m.group("neg1")) or m.group("sign") == "-" or (m.group("trail") or "").upper() in ("-", "DR")
    if negative:
        value = -value
    if not m.group("dec") and value.is_integer():
        return int(value)
    return value


# --------------------------------------------------------------------------
# Writing Excel
# --------------------------------------------------------------------------
def write_workbook(path: str, rows_by_page: Dict[int, List[List[str]]], options: Options,
                   one_sheet: bool = True, page_column: bool = True, source_name: str = ""):
    """Saves the extraction as .xlsx. One sheet with every page (optionally
    a "Page" column first) or one sheet per page. Returns the number of
    rows written."""
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Converted"
    total = 0

    def put(sheet, values, page_no=None):
        row = []
        if page_no is not None:
            row.append(page_no)
        for v in values:
            n = to_number(v) if options.convert_numbers else None
            row.append(n if n is not None else v)
        sheet.append(row)

    sheets = []
    if one_sheet:
        for i in sorted(rows_by_page):
            for r in rows_by_page[i]:
                put(ws, r, (i + 1) if page_column else None)
                total += 1
        sheets.append(ws)
    else:
        wb.remove(ws)
        for i in sorted(rows_by_page):
            s = wb.create_sheet(f"Page {i + 1}")
            for r in rows_by_page[i]:
                put(s, r)
                total += 1
            sheets.append(s)
        if not sheets:
            sheets.append(wb.create_sheet("Converted"))

    for s in sheets:
        widths = {}
        for row in s.iter_rows():
            for c in row:
                if c.value is not None:
                    widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
        for col, w in widths.items():
            s.column_dimensions[get_column_letter(col)].width = min(max(w + 2, 6), 60)
        s.freeze_panes = "A2" if s.max_row > 1 else None
        if s.max_row >= 1:
            for c in s[1]:
                c.font = Font(bold=True)
    info = wb.create_sheet("About")
    info.append(["Converted from", source_name])
    info.append(["Pages", ", ".join(str(i + 1) for i in sorted(rows_by_page))])
    info.append(["Note", "Check totals against the PDF before relying on this conversion."])
    wb.save(path)
    return total


def parse_page_range(text: str, page_count: int) -> List[int]:
    """'1-3, 5' -> [0, 1, 2, 4] (0-based). Blank = every page. Raises
    ValueError on anything it can't read or out of range."""
    text = (text or "").strip()
    if not text:
        return list(range(page_count))
    pages = []
    for part in text.split(","):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^(\d+)\s*(?:-\s*(\d+))?$", part)
        if not m:
            raise ValueError(f"Can't read the page range '{part}'. Use e.g. 1-3, 5")
        a = int(m.group(1)); b = int(m.group(2) or a)
        if a < 1 or b > page_count or a > b:
            raise ValueError(f"Page(s) {part} are outside this PDF (it has {page_count} page(s)).")
        pages.extend(range(a - 1, b))
    return sorted(set(pages))
