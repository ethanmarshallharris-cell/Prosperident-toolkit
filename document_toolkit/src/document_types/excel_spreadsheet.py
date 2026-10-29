"""
Excel Spreadsheet (choose columns) -- a document-type plugin, but a
different SHAPE of one than dentrix_audit_trail.py. See base.py's
"TWO SHAPES OF PLUGIN" section for why: there's no free-text pattern to
scan for here, so this plugin sets FORMAT_KIND = "spreadsheet" and
implements a different, smaller contract -- the user picks which
column(s) to redact by name, rather than the plugin auto-detecting
matches the way the PDF-report plugins do.

--------------------------------------------------------------------------
WHAT THIS PLUGIN DOES
--------------------------------------------------------------------------
Given an .xlsx workbook:
  1. list_columns() reads every worksheet's header row (row 1) and returns
     the distinct column names found, each with a sample value, so the GUI
     can show a checkbox per column.
  2. The user checks the column(s) to redact and, if they want, edits the
     placeholder label for each (default: the column's own name).
  3. apply_column_redaction() replaces every non-blank cell in each chosen
     column with a placeholder -- <Label 1>, <Label 2>, etc -- assigning
     the SAME placeholder to the SAME underlying value every time it
     recurs in that column (so, e.g., every row for the same SSN gets the
     same placeholder), independently per column (numbering restarts at 1
     in each column; a placeholder is only ever compared for reuse within
     the column it came from, never across columns).
  4. write_key() saves a CSV recording every (column, placeholder, real
     value) triple, so the Un-redact tab can reverse it later -- read by
     spreadsheet_unredact_engine.py, a sibling to unredact_engine.py (the
     PDF one). main_gui.py picks between them by the redacted file's
     extension.

A column is matched by its header text (row 1), independently per
worksheet -- a workbook with several sheets that all have (for instance)
a "Patient Name" column will have that column redacted on every sheet
that has it, not just the first. A sheet with none of the chosen column
names is left completely untouched.

--------------------------------------------------------------------------
COMBINED COLUMNS (added October 2026)
--------------------------------------------------------------------------
Two or more columns can be redacted together as ONE identity -- e.g. a
"First Name" column and a "Last Name" column redacted as <Name 1>,
<Name 2>, ... where each number stands for a distinct first+last
COMBINATION, not for a first name or a last name on its own. (Redacted
separately, "John Smith" and "John Doe" would share <First_Name 1> and
"Mary Smith" would share <Last_Name 1> with John Smith -- the numbering
would describe fragments, not people.)

A combined entry in `selected_columns` looks like
    {"name": combined_column_name(parts), "label": "Name", "combine": parts}
with `parts` the real header names in the user's chosen order. Every
part cell that has a value is overwritten with the SAME placeholder, so a
redacted row reads First Name = <Name 3>, Last Name = <Name 3>. A blank
part stays blank; a row where every part is blank is left alone.

The key file records, for each combination:
  - one row per part, under that part's own real column name
    ("First Name", <Name 3>, "John") and ("Last Name", <Name 3>, "Smith")
    -- exactly the ordinary row shape, so spreadsheet_unredact_engine.py
    restores each cell by its header with no special handling (and an
    older copy of the app can still un-redact the file); and
  - one summary row under the combined column name
    ("Combined: First Name + Last Name", <Name 3>, '["John", "Smith"]')
    -- the combination itself, JSON-encoded so a value containing spaces
    or punctuation can never be confused with a different split. This is
    what "Use existing key file(s)" and Batch Redact read to give the same
    person the same number in the next file. The un-redact engine skips
    these rows (they name no real column).

Identity: two rows are the same combination only if every part matches
exactly (same rule as a single column -- no case folding, no trimming),
because un-redaction must put back each cell's exact original text.

--------------------------------------------------------------------------
WHY THIS COUNTS AS TRUE REDACTION
--------------------------------------------------------------------------
Cell values are set directly (cell.value = placeholder_text) and the
workbook is then written out fresh via openpyxl -- the original value
is not hidden behind formatting or left in the file's shared-string table
for someone to recover; it is simply not present anywhere in the output
file's XML once saved. This is the same "genuinely removed, not covered
up" standard the PDF plugins hold themselves to, just achieved differently
because a spreadsheet cell (unlike a PDF content stream run) has no
in-place partial-overwrite ambiguity to begin with -- replacing a cell's
value replaces all of it, there's no analog to the tightly-spaced-PDF-line
problem that motivated the whole-page-rebuild strategy used elsewhere.

--------------------------------------------------------------------------
KNOWN LIMITATIONS
--------------------------------------------------------------------------
- Modern Excel (.xlsx) workbooks and .csv files are supported -- not the
  legacy .xls binary format. A .csv (added October 2026) is read through
  csv_table.py as a one-sheet workbook (sheet name "CSV"), so everything
  above -- choosing and combining columns, numbering, keys, one-stop keys,
  verification, batch redaction -- works on it unchanged. Every CSV value
  stays text exactly as written (no "00123" -> 123), and the redacted
  .csv is written back in the original's encoding, delimiter, quoting and
  line endings. A .csv is always saved as .csv.
- Only row 1 of each sheet is checked for headers. A workbook whose
  tables don't start at row 1 (title rows, merged banner cells, etc.
  above the real header) won't have its columns detected -- there is no
  "Needs review" concept here the way the PDF plugins have, since this
  plugin doesn't scan free text for candidate matches; either a column
  shows up in the checklist or it doesn't. This is a case where a
  redaction can be missed entirely, not just mis-parsed, precisely
  because it isn't found in the first place -- if a workbook's layout is
  unusual, check the Excel file's own header row before trusting the
  "Load Columns" list.
- Cell formatting (number format, fill color, conditional formatting,
  etc.) on a redacted cell is not specifically preserved or restored --
  openpyxl generally keeps a cell's existing style when only .value is
  changed, but formulas, charts, or other cells that referenced the
  original value are not rewritten and may show errors or stale results.
- Merged cells: only the top-left cell of a merged range carries a value
  in openpyxl's model, so a selected column that runs through merged
  cells will only have that anchor cell's value (and hence be reachable
  for redaction) -- the other cells in the merge have no value to find.
"""
import csv
import json
import os
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Dict, List

import openpyxl
from openpyxl.utils import get_column_letter

from . import _key_lineage
import csv_table

ID = "excel_spreadsheet"
DISPLAY_NAME = "Excel or CSV Spreadsheet (choose columns)"
FORMAT_KIND = "spreadsheet"
FILE_TYPES = [
    ("Excel or CSV files", "*.xlsx *.csv"),
    ("Excel files", "*.xlsx"),
    ("CSV files", "*.csv"),
    ("All files", "*.*"),
]
SUPPORTS_KEY_FILE = True
PROGRESS_UNIT_LABEL = "row"
# See base.py's "BATCH REDACTION" section and apply_column_redaction()'s
# `state` parameter below -- main_gui.py's Batch Redact tab checks this
# before offering batch mode for a document type.
SUPPORTS_BATCH = True

# The next three exist only so main_gui.py can read them unconditionally
# the same way it does for pdf_report-shaped plugins, without special
# casing every constant access by FORMAT_KIND. This plugin doesn't scan
# free text for candidate matches, so it has no real "entity" or "needs
# review" concept -- these are placeholders, not used to build any UI text
# that's actually shown for this document type.
ENTITY_LABEL = "Value"
CONTEXT_COLUMN_LABEL = ""
REVIEW_HINT = ""
NO_MATCHES_WARNING = (
    "No column headers were found in row 1 of any worksheet in this file. "
    "This plugin only looks at row 1 -- if this workbook's table starts "
    "somewhere else (a title row, merged banner cells, etc. above the real "
    "header), columns can't be auto-detected. Check the file's layout, or "
    "move the header row to row 1 and try again."
)


# Prefix of the synthetic column name a combined-columns entry uses in
# selected_columns, in `state`, and in the key file's summary rows. Must
# match spreadsheet_unredact_engine.COMBINED_COLUMN_PREFIX (kept as a
# separate constant there on purpose -- the two modules don't import each
# other).
COMBINED_COLUMN_PREFIX = "Combined: "
_COMBINED_PART_SEPARATOR = " + "


def combined_column_name(parts: List[str]) -> str:
    """The synthetic name for a combined-columns entry, e.g.
    "Combined: First Name + Last Name". Part order matters (it is the
    order values are stored in), so the same columns combined in a
    different order are a different entry."""
    return COMBINED_COLUMN_PREFIX + _COMBINED_PART_SEPARATOR.join(parts)


def is_combined_column_name(name) -> bool:
    return isinstance(name, str) and name.startswith(COMBINED_COLUMN_PREFIX)


def encode_combined_value(values) -> str:
    """One combination of part values -> the string stored as its "real
    value" in state and in the key file's summary row. JSON, so the split
    between parts is unambiguous whatever the values contain."""
    return json.dumps([str(v) for v in values], ensure_ascii=False)


def decode_combined_value(text: str) -> List[str]:
    values = json.loads(text)
    if not isinstance(values, list):
        raise ValueError("not a combined value")
    return [str(v) for v in values]


def combined_identity_text(values) -> str:
    """The text used to compare a combination against another document
    type's identities (cross-type "Same identity as:" linking): the
    non-blank parts joined with ", " in the order the user chose. With the
    last-name column first this reads "Smith, John" -- the way Dentrix
    writes patient names -- so the linked type's own normalize_identity()
    can match it. No reordering is ever guessed (see base.py)."""
    return ", ".join(str(v).strip() for v in values if str(v).strip())


def _split_selection(selected_columns: List[dict]):
    """(plain {name: label}, [combined entries], [full-name entries]) from
    a selected_columns list, after checking no real column is used twice --
    once on its own and once in a combination, or in two combinations.
    Either would make one cell receive two different placeholders, so it
    is refused. See "FULL-NAME COLUMNS" for the third kind."""
    plain = {}
    groups = []
    fullnames = []
    used = {}
    for entry in selected_columns:
        fn = entry.get("full_name")
        if fn:
            name = entry["name"]
            if name in used:
                raise ValueError(f'Column "{name}" is selected more than once ({used[name]} and as a full name).')
            used[name] = "as a full name"
            parts = [str(p) for p in fn["parts"]]
            if fn["first"] not in parts or fn["last"] not in parts or fn["first"] == fn["last"]:
                raise ValueError(f"Full-name column '{name}': first and last name must be two of {parts}.")
            if fn.get("format") not in NAME_FORMATS:
                raise ValueError(f"Full-name column '{name}': unknown name format {fn.get('format')!r}.")
            fullnames.append({
                "name": name, "label": entry["label"], "combined": fn["combined"], "parts": parts,
                "first": fn["first"], "last": fn["last"], "format": fn["format"],
            })
            continue
        parts = entry.get("combine")
        if parts:
            parts = [str(p) for p in parts]
            if len(parts) < 2:
                raise ValueError(f"A combined entry needs at least two columns: {parts}")
            if len(set(parts)) != len(parts):
                raise ValueError(f"The same column is listed twice in one combination: {parts}")
            owner = entry.get("name") or combined_column_name(parts)
            for p in parts:
                if p in used:
                    raise ValueError(f'Column "{p}" is selected more than once ({used[p]} and {owner}).')
                used[p] = owner
            groups.append({"name": owner, "label": entry["label"], "parts": parts})
        else:
            name = entry["name"]
            if name in used:
                raise ValueError(f'Column "{name}" is selected more than once ({used[name]} and on its own).')
            used[name] = "on its own"
            plain[name] = entry["label"]
    return plain, groups, fullnames


# --------------------------------------------------------------------------
# FULL-NAME COLUMNS (added October 2026)
# --------------------------------------------------------------------------
# One column holding a whole name ("Smith, John") numbered the SAME as an
# existing First Name + Last Name combination -- from a key file loaded
# with "Use existing key file(s)", or a combination in the same batch --
# so the same person gets the same <Name N> whichever way a file writes
# the name. An entry in selected_columns:
#     {"name": "Patient", "label": "Name",
#      "full_name": {"combined": "Combined: First Name + Last Name",
#                    "parts": ["First Name", "Last Name"],
#                    "first": "First Name", "last": "Last Name",
#                    "format": "last_first"}}
# Formats: "last_first" (Smith, John), "first_last" (John Smith),
# "last_first_space" (Smith John). A comma always means "Last, First",
# whichever format is chosen.
#
# MATCHING, in order (first that gives exactly ONE person wins):
#   1. the exact cell text, if an earlier full-name redaction recorded it;
#   2. same first and last name, ignoring capitals, periods and extra
#      spaces ("SMITH, JOHN" = John Smith). Without a comma, the whole
#      name is also compared word for word, so "Mary Ann Lee" can still
#      match First "Mary Ann" + Last "Lee";
#   3. APPROXIMATE: same last name and same first word of the first name
#      ("Smith, John A." = John Smith), only when exactly one person fits.
# Two or more people fitting is AMBIGUOUS: the cell gets a new number of
# its own -- never someone else's -- and is listed for checking. A name
# that matches nobody gets the next new number in the combination, and
# the key records it as a combination too (first/last parts), so later
# files written either way match it.
#
# The key also gets an ordinary row (this column, placeholder, the cell's
# exact text) so un-redaction restores the column exactly. If a file
# spells one person two ways ("Smith, John" and "SMITH, JOHN") both get
# the same number, but un-redaction can only restore one spelling: the
# first one is kept and the others are listed in the summary.
NAME_FORMATS = {
    "last_first": "Last, First  (e.g. Smith, John)",
    "first_last": "First Last  (e.g. John Smith)",
    "last_first_space": "Last First, no comma  (e.g. Smith John)",
}

_PLACEHOLDER_NUMBER_RE = re.compile(r"^<(.+) (\d+)>$")


def _norm_name(text) -> str:
    return " ".join(str(text).replace(",", " ").replace(".", " ").lower().split())


def parse_full_name(text: str, fmt: str):
    """(first, last) from a full-name cell, per the chosen format (a comma
    always splits "Last, First")."""
    t = str(text).strip()
    if "," in t:
        last, first = t.split(",", 1)
        return first.strip(), last.strip()
    words = t.split()
    if not words:
        return "", ""
    if fmt == "first_last":
        return " ".join(words[:-1]), words[-1]
    return " ".join(words[1:]), words[0]


def placeholder_label_and_number(placeholder):
    m = _PLACEHOLDER_NUMBER_RE.match(str(placeholder))
    return (m.group(1), int(m.group(2))) if m else (None, None)


class _FullNameMatcher:
    """Looks up / allocates the combination number for full-name cells
    against one combination's {encoded: placeholder} map (shared with the
    combination itself when it is redacted in the same file)."""

    def __init__(self, group_map, parts, first, last, label, counter_ref):
        self.map = group_map
        self.parts = parts
        self.fi = parts.index(first)
        self.li = parts.index(last)
        self.label = label
        self.counter = counter_ref        # {"n": next-number-minus-one}, shared
        self.pair = {}                    # (first, last) normalized -> {placeholder}
        self.seq_fl = {}                  # "first last" -> {placeholder}
        self.seq_lf = {}                  # "last first" -> {placeholder}
        self.first_word = {}              # (first word, last) -> {placeholder}
        self.display = {}                 # placeholder -> "First Last" (for reports)
        self.siblings = [self]            # every matcher sharing this map (several
                                          # full-name columns -> one combination)
        for encoded, ph in list(group_map.items()):
            self._index(encoded, ph)

    def _index(self, encoded, ph):
        try:
            values = decode_combined_value(encoded)
        except (ValueError, TypeError):
            return
        if len(values) != len(self.parts):
            return
        f, l = _norm_name(values[self.fi]), _norm_name(values[self.li])
        if not f and not l:
            return
        self.pair.setdefault((f, l), set()).add(ph)
        self.seq_fl.setdefault(f"{f} {l}".strip(), set()).add(ph)
        self.seq_lf.setdefault(f"{l} {f}".strip(), set()).add(ph)
        if f:
            self.first_word.setdefault((f.split()[0], l), set()).add(ph)
        self.display.setdefault(ph, f"{values[self.fi]} {values[self.li]}".strip())

    def resolve(self, text, fmt):
        """(placeholder, how) -- how is "exact", "approximate", "ambiguous" or "new"."""
        first, last = parse_full_name(text, fmt)
        f, l = _norm_name(first), _norm_name(last)
        hits = set(self.pair.get((f, l), set()))
        if "," not in text:
            whole = _norm_name(text)
            hits |= (self.seq_fl if fmt == "first_last" else self.seq_lf).get(whole, set())
        if len(hits) == 1:
            return next(iter(hits)), "exact"
        how = "ambiguous" if hits else None
        if not hits and f:
            approx = self.first_word.get((f.split()[0], l), set())
            if len(approx) == 1:
                return next(iter(approx)), "approximate"
            if len(approx) > 1:
                how = "ambiguous"
        values = [""] * len(self.parts)
        values[self.fi], values[self.li] = first, last
        encoded = encode_combined_value(values)
        if encoded in self.map and how is None:
            return self.map[encoded], "exact"
        if how == "ambiguous" or encoded in self.map:
            # Never give this cell another person's number: a number of
            # its own, recorded under the cell's own text only.
            self.counter["n"] += 1
            return f"<{self.label} {self.counter['n']}>", "ambiguous"
        self.counter["n"] += 1
        ph = f"<{self.label} {self.counter['n']}>"
        self.map[encoded] = ph
        for m in self.siblings:
            m._index(encoded, ph)
        return ph, "new"


def _combined_key_rows(group_name: str, parts: List[str], value_map: Dict[str, str]) -> List[dict]:
    """Key rows for one combined entry: a summary row per combination plus
    one ordinary row per non-blank part (see module docstring)."""
    rows = []
    for encoded, placeholder in value_map.items():
        rows.append({"column": group_name, "placeholder": placeholder, "value": encoded})
        try:
            values = decode_combined_value(encoded)
        except (ValueError, TypeError):
            continue
        for part, value in zip(parts, values):
            if value.strip():
                rows.append({"column": part, "placeholder": placeholder, "value": value})
    return rows


@dataclass
class ColumnInfo:
    """One distinct column header found across the workbook's sheet(s)."""
    name: str               # exact header text as found in row 1
    sheets: List[str]       # worksheet name(s) this header appears in
    sample: str              # first non-blank cell value seen under this header, or ""
    default_label: str       # sanitized default placeholder label, e.g. "Patient_Name"


def _sanitize_label(name: str) -> str:
    """Turns a column header into a safe placeholder label: whitespace runs
    collapse to a single underscore, anything that isn't a letter, digit,
    underscore or hyphen is dropped. Falls back to "Column" if that leaves
    nothing usable (e.g. a header that was purely punctuation/emoji)."""
    collapsed = re.sub(r"\s+", "_", name.strip())
    cleaned = re.sub(r"[^A-Za-z0-9_\-]", "", collapsed)
    return cleaned or "Column"


def _open_workbook(path: str, **kwargs):
    """openpyxl.load_workbook(), but with a much more actionable error if
    it fails. In practice, a workbook that openpyxl can't parse is almost
    always one that wasn't saved by genuine Microsoft Excel -- exported
    from a practice-management system, general ledger, or other reporting
    tool -- and has some structural quirk in its internal XML (a missing
    relationship id, an out-of-range style reference, etc.) that Excel
    itself silently tolerates but openpyxl's stricter parser doesn't. The
    single most reliable fix for this, confirmed to work across a wide
    range of these quirks, is to open the file in genuine Excel and use
    File > Save As to write a fresh copy -- Excel rewrites the internal
    XML cleanly when it does this, which clears out whatever the original
    exporting program left broken. Without the actual failing file in
    hand, guessing which specific quirk this run hit isn't reliable, so
    this re-raises with the original exception type/message preserved
    (via `raise ... from exc`) rather than trying to interpret it.
    """
    if csv_table.is_csv(path):
        return csv_table.load_as_workbook(path)
    try:
        return openpyxl.load_workbook(path, **kwargs)
    except Exception as exc:
        raise RuntimeError(
            f"openpyxl could not open this workbook ({type(exc).__name__}: {exc}). "
            "This usually means the file wasn't saved by genuine Microsoft Excel (e.g. "
            "it was exported from another program) and has some structural quirk in its "
            "internal XML that Excel tolerates but openpyxl doesn't. The most reliable "
            "fix: open the file in Excel and use File > Save As to save a fresh copy in "
            ".xlsx format, then use that copy instead."
        ) from exc


def list_columns(path: str) -> List[ColumnInfo]:
    """Opens the workbook read-only and returns every distinct column
    header found in row 1 of any worksheet, in first-seen order. Does not
    modify anything.
    """
    wb = _open_workbook(path, read_only=True, data_only=True)
    try:
        columns: "Dict[str, ColumnInfo]" = {}
        for ws in wb.worksheets:
            rows = ws.iter_rows()
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            header_by_idx = {}
            for cell in header_row:
                if cell.value is None:
                    continue
                text = str(cell.value).strip()
                if not text:
                    continue
                header_by_idx[cell.column] = text
                if text not in columns:
                    columns[text] = ColumnInfo(
                        name=text, sheets=[], sample="", default_label=_sanitize_label(text)
                    )
                if ws.title not in columns[text].sheets:
                    columns[text].sheets.append(ws.title)
            if not header_by_idx:
                continue
            # Grab one sample value per header, from the first data row that has one.
            still_needed = {idx for idx in header_by_idx if not columns[header_by_idx[idx]].sample}
            if not still_needed:
                continue
            for row in rows:
                if not still_needed:
                    break
                for cell in row:
                    # Check .value BEFORE .column, always, in this order --
                    # read_only mode pads a sparse row (any blank cell
                    # sitting between two filled columns, extremely common
                    # in real spreadsheets) with openpyxl's EmptyCell
                    # placeholder, which has .value (always None) but has
                    # NO .column attribute at all. Reading .column first,
                    # as an earlier version of this loop did, raised
                    # AttributeError the moment a sparse row was hit.
                    if cell.value in (None, ""):
                        continue
                    if cell.column in still_needed:
                        col = columns[header_by_idx[cell.column]]
                        col.sample = str(cell.value)
                        still_needed.discard(cell.column)
        return list(columns.values())
    finally:
        wb.close()


def distinct_values_in_column(path: str, column_name: str) -> "set[str]":
    """Opens the workbook read-only and returns every distinct non-blank
    cell value (as str) found under the given header, across every sheet
    that has it in row 1. Read-only, single-purpose scan -- does not
    modify anything and does not build the full placeholder maps that
    apply_column_redaction() does.

    Used by the Batch Redact tab's cross-type identity linking (see
    base.py's "BATCH REDACTION" section): before redacting a file whose
    column is manually linked to another document type's identity (e.g. a
    "Patient Name" column linked to the Dentrix plugin's Patient
    numbering), the caller needs every distinct value that will appear in
    that column so each one can be resolved against the shared identity
    pool ahead of time -- this is what supplies that list.
    """
    wb = _open_workbook(path, read_only=True, data_only=True)
    try:
        values = set()
        for ws in wb.worksheets:
            rows = ws.iter_rows()
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            target_idx = None
            for cell in header_row:
                if cell.value is not None and str(cell.value).strip() == column_name:
                    target_idx = cell.column
                    break
            if target_idx is None:
                continue
            for row in rows:
                for cell in row:
                    # See list_columns()'s comment on this same pattern --
                    # check .value before .column, a sparse row's padding
                    # cells have no .column attribute at all.
                    if cell.value in (None, ""):
                        continue
                    if cell.column == target_idx:
                        values.add(str(cell.value))
        return values
    finally:
        wb.close()


def distinct_combined_values(path: str, parts: List[str]) -> "set[tuple]":
    """Every distinct combination of the given columns' values (a tuple
    of str, "" for a blank part) across every sheet that has at least one
    of them in row 1. Rows where every part is blank are skipped. Read-only.

    The combined-columns counterpart of distinct_values_in_column(), for
    cross-type identity linking of a combination (e.g. First + Last Name
    linked to the Dentrix "Patient" identity)."""
    wb = _open_workbook(path, read_only=True, data_only=True)
    try:
        combos = set()
        for ws in wb.worksheets:
            rows = ws.iter_rows()
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            idx_by_part = {}
            for cell in header_row:
                if cell.value is None:
                    continue
                text = str(cell.value).strip()
                if text in parts and text not in idx_by_part:
                    idx_by_part[text] = cell.column
            if not idx_by_part:
                continue
            for row in rows:
                by_col = {}
                for cell in row:
                    # .value before .column -- see list_columns().
                    if cell.value in (None, ""):
                        continue
                    by_col[cell.column] = cell.value
                values = []
                for p in parts:
                    v = by_col.get(idx_by_part.get(p))
                    values.append("" if v is None or str(v).strip() == "" else str(v))
                if any(v for v in values):
                    combos.add(tuple(values))
        return combos
    finally:
        wb.close()


def _resolve_full_name(col_name, text, col_map, matcher, fmt, cache, report):
    """The placeholder for one full-name cell (see "FULL-NAME COLUMNS"),
    recording how it was matched in `report`. `col_map` is the column's
    {cell text: placeholder} map that becomes its key rows: only the FIRST
    spelling of each person goes in, so un-redaction is unambiguous."""
    key = (col_name, text)
    if key in cache:
        return cache[key]
    ph = col_map.get(text)
    how = "exact" if ph else None
    if ph is None:
        ph, how = matcher.resolve(text, fmt)
        if ph in col_map.values():
            first = next(v for v, p in col_map.items() if p == ph)
            report["variants"][text] = f"{ph} (restores as {first!r})"
        else:
            col_map[text] = ph
    if how == "approximate":
        report["approximate"][text] = f"{ph} = {matcher.display.get(ph, '')}"
    elif how == "ambiguous":
        report["ambiguous"][text] = ph
    elif how == "new":
        report["new"].add(ph)
    else:
        report["matched"].add(ph)
    cache[key] = ph
    return ph


def apply_column_redaction(
    input_path: str,
    output_path: str,
    selected_columns: List[dict],
    progress_callback=None,
    phase_callback=None,
    state=None,
) -> "tuple[dict, list]":
    """Writes a copy of input_path with every non-blank cell in each chosen
    column replaced by a placeholder. `selected_columns` is
    [{"name": <exact header text>, "label": <placeholder label to use>}, ...]
    -- `label` is what goes inside the angle brackets; the caller (the GUI)
    is responsible for defaulting it to the column's own sanitized name and
    letting the user override it. An entry with a "combine" list redacts
    those columns together as one identity -- see the module docstring's
    "COMBINED COLUMNS" section. Its `state` entry is keyed by its synthetic
    combined_column_name(), with encode_combined_value() strings as values.

    The same real value always gets the same placeholder within its own
    column (independent numbering per column -- see module docstring).

    progress_callback(current, total) is called with 1-based row progress
    across every sheet that has at least one chosen column, if given; any
    exception it raises is swallowed. phase_callback(phase_name) is called
    once, with "saving", right before the final write to disk.

    `state`, if given, is a running {column_name: {value: placeholder_text}}
    mapping to seed this call's numbering from, instead of starting every
    column's counter fresh at 1 -- see base.py's "BATCH REDACTION" section.
    Passing the SAME dict through a series of calls (one per file in a
    batch) makes a value that recurs across those files keep the same
    placeholder text everywhere, rather than each file numbering
    independently and risking the same placeholder text meaning two
    different values in two different files. This function does not
    mutate `state` itself (a fresh dict is built from it internally) --
    the caller is expected to fold each call's returned key_rows back into
    its own running state before the next call; see main_gui.py's batch
    redact worker for the actual accumulation loop. Omit (or pass None)
    for the original single-file behavior, unchanged.

    Returns (summary, key_rows):
      summary: {'columns_selected': int, 'columns_found': int,
                'columns_not_found': list[str], 'cells_redacted': int,
                'unique_values_redacted': int}
      key_rows: [{'column': str, 'placeholder': str, 'value': str}, ...],
                one row per unique (column, value) pair KNOWN SO FAR for a
                column found in this file -- including any carried forward
                via `state` even if this particular file didn't happen to
                contain that exact value -- in the order placeholders were
                assigned. This is what write_key() saves for a single file;
                a batch caller should accumulate across every call's
                key_rows rather than assume the last file's alone covers
                the whole batch (a column absent from the last file
                wouldn't otherwise be represented).
    """
    label_by_name, groups, fullnames = _split_selection(selected_columns)
    # A full-name column is found and seeded like an ordinary column (its
    # map holds the cell texts already numbered); its cells are numbered
    # through the combination instead -- see "FULL-NAME COLUMNS".
    fullname_by_col = {fn["name"]: fn for fn in fullnames}
    for fn in fullnames:
        label_by_name[fn["name"]] = fn["label"]
    all_parts = {p for g in groups for p in g["parts"]}
    wb = _open_workbook(input_path)
    try:
        # Pass 1: find which sheets have which chosen columns, and the
        # total row count across them, so progress can be reported.
        sheet_col_map = {}   # ws.title -> {col_idx: header_name}
        sheet_group_map = {}  # ws.title -> [(group, {part: col_idx or None}), ...]
        found_names = set()
        found_groups = set()
        parts_missing = []    # "Last Name (sheet 'Visits')" -- a combination found on a
                              # sheet that lacks one of its parts (treated as blank there)
        for ws in wb.worksheets:
            col_idx_by_name = {}
            idx_by_header = {}
            try:
                header_row = next(ws.iter_rows(min_row=1, max_row=1))
            except StopIteration:
                continue
            for cell in header_row:
                if cell.value is None:
                    continue
                text = str(cell.value).strip()
                if text in label_by_name:
                    col_idx_by_name[cell.column] = text
                    found_names.add(text)
                if text in all_parts and text not in idx_by_header:
                    idx_by_header[text] = cell.column
            if col_idx_by_name:
                sheet_col_map[ws.title] = col_idx_by_name
            for g in groups:
                idx_by_part = {p: idx_by_header.get(p) for p in g["parts"]}
                if any(i is not None for i in idx_by_part.values()):
                    sheet_group_map.setdefault(ws.title, []).append((g, idx_by_part))
                    found_groups.add(g["name"])
                    for p, i in idx_by_part.items():
                        if i is None:
                            parts_missing.append(f"{p} (sheet '{ws.title}')")

        redacted_titles = list(sheet_col_map)
        redacted_titles += [t for t in sheet_group_map if t not in sheet_col_map]
        total_rows = sum(
            max(wb[title].max_row - 1, 0) for title in redacted_titles
        )

        state = state or {}
        # Seed each found column's map from any carried-forward state (see
        # this function's docstring) -- for a fresh single-file call,
        # state is {} so this is identical to the old `{name: {} ...}`.
        placeholder_maps: "Dict[str, Dict[str, str]]" = {
            name: dict(state.get(name, {})) for name in found_names
        }
        # A column's counter is just how many placeholders it already has --
        # correct whether those came from this file or were carried forward,
        # since placeholders are always assigned as a gapless 1..N sequence.
        counters = {name: len(placeholder_maps[name]) for name in found_names}
        group_maps: "Dict[str, Dict[str, str]]" = {
            g["name"]: dict(state.get(g["name"], {})) for g in groups if g["name"] in found_groups
        }
        # Next number per combination: past every number already used for
        # its label anywhere in the carried-forward state (not just the
        # combination's own count), so a number given to an ambiguous
        # full-name cell in an earlier file is never handed out again.
        def _start_number(name, label):
            used = [placeholder_label_and_number(ph) for ph in group_maps[name].values()]
            labels = {lab for lab, _n in used if lab} or {label}
            top = len(group_maps[name])
            for col_map in state.values():
                if not isinstance(col_map, dict):
                    continue
                for ph in col_map.values():
                    lab, n = placeholder_label_and_number(ph)
                    if lab in labels and n > top:
                        top = n
            return top

        extra_groups = []       # combinations numbered only through a full-name column
        fullname_matchers = {}  # full-name column -> _FullNameMatcher
        fullname_report = {}    # full-name column -> counts / examples for the summary
        for fn in fullnames:
            if fn["name"] not in found_names:
                continue
            cname = fn["combined"]
            if cname not in group_maps:
                group_maps[cname] = dict(state.get(cname, {}))
                extra_groups.append({"name": cname, "parts": fn["parts"], "label": fn["label"]})
        label_of_group = {g["name"]: g["label"] for g in groups + extra_groups}
        group_counters = {name: {"n": _start_number(name, label_of_group[name])} for name in group_maps}
        for fn in fullnames:
            if fn["name"] not in found_names:
                continue
            cname = fn["combined"]
            existing_labels = [placeholder_label_and_number(ph)[0] for ph in group_maps[cname].values()]
            label = label_of_group.get(cname) if cname not in {g["name"] for g in extra_groups} else None
            label = label or next((lab for lab in existing_labels if lab), None) or fn["label"]
            fullname_matchers[fn["name"]] = _FullNameMatcher(
                group_maps[cname], fn["parts"], fn["first"], fn["last"], label, group_counters[cname]
            )
            for other in fullname_matchers.values():
                if other.map is group_maps[cname] and other is not fullname_matchers[fn["name"]]:
                    other.siblings.append(fullname_matchers[fn["name"]])
                    fullname_matchers[fn["name"]].siblings = other.siblings
                    break
            fullname_report[fn["name"]] = {
                "combination": cname, "matched": set(), "new": set(), "approximate": {},
                "ambiguous": {}, "variants": {},
            }
        fullname_cache = {}     # (column, cell text) -> placeholder
        cells_redacted = 0
        current = 0

        for title in redacted_titles:
            col_idx_by_name = sheet_col_map.get(title, {})
            sheet_groups = sheet_group_map.get(title, [])
            ws = wb[title]
            for row in ws.iter_rows(min_row=2):
                current += 1
                if sheet_groups:
                    cell_by_col = {}
                    for cell in row:
                        col = getattr(cell, "column", None)
                        if col is not None:
                            cell_by_col[col] = cell
                    for g, idx_by_part in sheet_groups:
                        part_cells = [cell_by_col.get(i) if i is not None else None for i in idx_by_part.values()]
                        values = [
                            "" if c is None or c.value is None or str(c.value).strip() == "" else str(c.value)
                            for c in part_cells
                        ]
                        if not any(values):
                            continue
                        encoded = encode_combined_value(values)
                        g_map = group_maps[g["name"]]
                        placeholder_text = g_map.get(encoded)
                        if placeholder_text is None:
                            group_counters[g["name"]]["n"] += 1
                            placeholder_text = f"<{g['label']} {group_counters[g['name']]['n']}>"
                            g_map[encoded] = placeholder_text
                            matcher_for = [m for m in fullname_matchers.values() if m.map is g_map]
                            for m in matcher_for:
                                m._index(encoded, placeholder_text)
                        for c, v in zip(part_cells, values):
                            if c is not None and v:
                                c.value = placeholder_text
                                cells_redacted += 1
                for cell in row:
                    col_name = col_idx_by_name.get(cell.column)
                    if col_name is None:
                        continue
                    if cell.value is None or str(cell.value).strip() == "":
                        continue
                    value_str = str(cell.value)
                    col_map = placeholder_maps[col_name]
                    if col_name in fullname_matchers:
                        cell.value = _resolve_full_name(
                            col_name, value_str, col_map, fullname_matchers[col_name],
                            fullname_by_col[col_name]["format"], fullname_cache, fullname_report[col_name],
                        )
                        cells_redacted += 1
                        continue
                    placeholder_text = col_map.get(value_str)
                    if placeholder_text is None:
                        counters[col_name] += 1
                        placeholder_text = f"<{label_by_name[col_name]} {counters[col_name]}>"
                        col_map[value_str] = placeholder_text
                    cell.value = placeholder_text
                    cells_redacted += 1
                if progress_callback is not None:
                    try:
                        progress_callback(current, max(total_rows, 1))
                    except Exception:
                        pass

        if phase_callback is not None:
            try:
                phase_callback("saving")
            except Exception:
                pass
        csv_table.save_workbook(wb, output_path)

        key_rows = []
        for col_name, col_map in placeholder_maps.items():
            for value_str, placeholder_text in col_map.items():
                key_rows.append({"column": col_name, "placeholder": placeholder_text, "value": value_str})
        for g in groups + extra_groups:
            if g["name"] in group_maps:
                key_rows.extend(_combined_key_rows(g["name"], g["parts"], group_maps[g["name"]]))

        not_found = sorted(set(label_by_name) - found_names)
        not_found += [g["name"] for g in groups if g["name"] not in found_groups]
        summary = {
            "columns_selected": len(selected_columns),
            "columns_found": len(found_names) + len(found_groups),
            "columns_not_found": not_found,
            "combined_parts_missing": parts_missing,
            "full_name_columns": {
                col: {
                    "combination": r["combination"],
                    "matched": len(r["matched"]),
                    "new": len(r["new"]),
                    "approximate": sorted(r["approximate"].items()),
                    "ambiguous": sorted(r["ambiguous"].items()),
                    "variants": sorted(r["variants"].items()),
                }
                for col, r in fullname_report.items()
            },
            "cells_redacted": cells_redacted,
            "unique_values_redacted": (
                sum(len(m) for m in placeholder_maps.values())
                + sum(len(m) for m in group_maps.values())
            ),
        }
        return summary, key_rows
    finally:
        wb.close()


def _scan_output_for_leaks_and_shape(output_path: str, placeholders_by_value: "Dict[str, List[dict]]"):
    """One read-only pass over output_path collecting both structured leak
    entries and each sheet's (max_row, max_column) shape -- both need a
    full walk of every cell, and there's no reason to make two passes over
    what can be a very large workbook (some real-world exports run to
    tens of thousands of rows) just to get them separately.

    Returns (leaks, out_shape). Each leak entry:
      {'sheet': str, 'coordinate': str, 'row': int, 'value': str,
       'is_header_row': bool, 'column_header': str or None,
       'row_context': [{'label': str, 'value': str}, ...],
       'placeholders': [{'column': str, 'placeholder': str}, ...]}
    'placeholders' lists every (column, placeholder) pair already on
    record for this exact value -- there can be more than one if the same
    value was redacted out of more than one selected column.

    'column_header' is the header text actually sitting in row 1 of the
    leak's OWN column (not necessarily one of the redacted columns listed
    in 'placeholders' -- a leak can just as easily turn up in an unrelated
    column, e.g. a free-text "Notes" field that happens to repeat a
    patient's name). 'row_context' is up to 4 other non-blank cell values
    from the SAME row (excluding the leak cell itself), each paired with
    its own column's header when known -- this is what lets a reviewer
    tell which actual record a leak belongs to (e.g. "Patient ID: 4521,
    Visit Date: 07/01/2016") without having to go open the workbook
    itself. Left empty for a header-row leak, where "other headers in
    this row" wouldn't mean anything as record context.
    """
    real_values = set(placeholders_by_value)
    leaks = []
    out_shape = {}
    wb_out = _open_workbook(output_path, read_only=True, data_only=True)
    try:
        for ws in wb_out.worksheets:
            out_shape[ws.title] = (ws.max_row, ws.max_column)
            if not real_values:
                continue
            headers_by_col = {}
            for row_idx, row in enumerate(ws.iter_rows(), start=1):
                if row_idx == 1:
                    # Populated before any leak in this row is processed
                    # below (same row, same iteration) -- and stays
                    # populated for every later row, since row 1 is always
                    # this loop's first iteration.
                    for cell in row:
                        if cell.value is not None:
                            text = str(cell.value).strip()
                            if text:
                                headers_by_col[cell.column] = text
                for cell in row:
                    if cell.value is None:
                        continue
                    val_str = str(cell.value)
                    if val_str not in real_values:
                        continue
                    is_header_row = row_idx == 1
                    row_context = []
                    if not is_header_row:
                        for other in row:
                            if other.column == cell.column or other.value in (None, ""):
                                continue
                            label = headers_by_col.get(other.column) or f"column {get_column_letter(other.column)}"
                            row_context.append({"label": label, "value": str(other.value)})
                            if len(row_context) >= 4:
                                break
                    leaks.append({
                        "sheet": ws.title,
                        "coordinate": cell.coordinate,
                        "row": cell.row,
                        "value": val_str,
                        # Row 1 is always the header row this plugin reads
                        # column names from, and redaction never touches it
                        # (apply_column_redaction() starts at min_row=2) --
                        # flagging this explicitly is what lets a
                        # header-label coincidence be told apart from a
                        # real leak at a glance, without needing to open
                        # the file and check.
                        "is_header_row": is_header_row,
                        "column_header": headers_by_col.get(cell.column),
                        "row_context": row_context,
                        "placeholders": placeholders_by_value[val_str],
                    })
    finally:
        wb_out.close()
    return leaks, out_shape


def verify_column_redaction(input_path: str, output_path: str, selected_columns: List[dict], key_rows: List[dict]) -> dict:
    """Independent post-check, same spirit as a pdf_report plugin's
    verify_redaction(): none of the real values that were redacted should
    still be found anywhere in the output workbook (not just in their own
    column -- a value that leaked into an unrelated cell is just as much a
    problem), and the workbook's overall shape (sheet names, sheet
    dimensions) should be unchanged, since only cell VALUES should have
    been touched.

    Returns {'ok': bool, 'issues': list[str], 'leaks': list[dict],
    'shape_issue': str or None, 'other_issue': str or None}.

    'issues' is the original human-readable summary (what the message-box
    error dialogs show). 'leaks' is the full, structured, UNtruncated list
    behind it (see _scan_output_for_leaks_and_shape()'s docstring for the
    shape of each entry) -- main_gui.py's manual-review screen acts on
    this directly, one leak at a time, rather than parsing 'issues'
    strings.

    'shape_issue' and 'other_issue', when set (else None), flag problems
    that are NOT a per-cell value leak a manual review screen can
    meaningfully resolve -- the workbook's own shape changed unexpectedly,
    or the source file couldn't even be re-opened to compare against.
    main_gui.py only offers manual review when 'leaks' is non-empty AND
    both of these are None; otherwise it falls back to the original hard
    block (delete the output, show an error, nothing saved).
    """
    issues = []
    other_issue = None
    shape_issue = None

    # Map each real value string to every (column, placeholder) pair
    # already on record for it, so a leaked cell can say WHICH selected
    # column's data this is -- e.g. "matches a value redacted from column
    # 'SSN'" -- and, for the manual-review screen, exactly which
    # placeholder text is available to replace it with. This adds no new
    # sensitive information to a message (a column header name and a
    # placeholder tag, not the value itself).
    placeholders_by_value: "Dict[str, List[dict]]" = {}
    for row in key_rows:
        if is_combined_column_name(row["column"]):
            # A combination's summary row -- its "value" is the JSON
            # encoding, never a cell's text. Its parts are checked through
            # their own ordinary rows.
            continue
        placeholders_by_value.setdefault(row["value"], []).append(
            {"column": row["column"], "placeholder": row["placeholder"]}
        )

    try:
        wb_in = _open_workbook(input_path, read_only=True, data_only=True)
        in_shape = {ws.title: (ws.max_row, ws.max_column) for ws in wb_in.worksheets}
        wb_in.close()
    except Exception as exc:  # noqa: BLE001
        other_issue = f"Could not re-open the source file to compare structure: {exc}"
        issues.append(other_issue)
        in_shape = None

    leaks, out_shape = _scan_output_for_leaks_and_shape(output_path, placeholders_by_value)

    if leaks:
        shown = []
        for leak in leaks[:10]:
            cols = ", ".join(sorted({p["column"] for p in leak["placeholders"]}))
            row_note = (
                " [row 1 -- this workbook's header row, not a data row]"
                if leak["is_header_row"] else ""
            )
            found_in = f" -- sitting in column '{leak['column_header']}'" if leak["column_header"] else ""
            shown.append(
                f"{leak['sheet']}!{leak['coordinate']}{row_note}{found_in} "
                f"(matches a value redacted from column '{cols}')"
            )
        more = f" (+{len(leaks) - 10} more)" if len(leaks) > 10 else ""
        issues.append(f"Original value(s) still present in output at: {'; '.join(shown)}{more}")

    if in_shape is not None and in_shape != out_shape:
        shape_issue = (
            "The workbook's sheet names or dimensions changed between input and output "
            "-- something beyond the chosen columns' values may have been affected "
            f"(before: {in_shape}, after: {out_shape})."
        )
        issues.append(shape_issue)

    return {
        "ok": len(issues) == 0,
        "issues": issues,
        "leaks": leaks,
        "shape_issue": shape_issue,
        "other_issue": other_issue,
    }


def apply_leak_decisions(output_path: str, decisions: List[dict]) -> int:
    """Applies manual-review decisions (see verify_column_redaction()'s
    'leaks' field) directly to output_path and re-saves it.

    decisions: [{'sheet': str, 'coordinate': str, 'action': 'replace' or
    'ignore', 'placeholder': str (required when action == 'replace')},
    ...] -- one entry per leak the review screen showed.

    A 'replace' decision overwrites that exact cell with the chosen
    placeholder text, the same as if the original redaction pass had
    caught it. An 'ignore' decision is a no-op here -- this function has
    no way to tell "reviewed and deliberately kept" apart from "never
    looked at", so the caller (main_gui.py) is responsible for recording
    every decision, including ignores, somewhere durable (the manual
    review log) before treating the file as finished.

    Returns the number of cells actually changed. The caller should
    re-run verify_column_redaction() afterward to confirm every leak was
    actually accounted for (replaced, or explicitly, recordedly ignored)
    rather than silently missed because of a coordinate mismatch upstream
    -- this function does not re-verify on its own.
    """
    to_replace = {
        (d["sheet"], d["coordinate"]): d["placeholder"]
        for d in decisions
        if d.get("action") == "replace"
    }
    if not to_replace:
        return 0
    changed = 0
    wb = _open_workbook(output_path)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    key = (ws.title, cell.coordinate)
                    if key in to_replace:
                        cell.value = to_replace[key]
                        changed += 1
        csv_table.save_workbook(wb, output_path)
    finally:
        wb.close()
    return changed


def _write_key_body(writer, key_rows: List[dict]) -> None:
    writer.writerow(["Column", "Placeholder", "Real Value"])
    for row in key_rows:
        writer.writerow([row["column"], row["placeholder"], row["value"]])


def write_key(key_path: str, input_path: str, output_path: str, key_rows: List[dict], carried_from=None) -> None:
    """Writes a CSV recording every (column, placeholder, real value)
    triple, so the redaction can be reversed later via the Un-redact tab
    (spreadsheet_unredact_engine.py). One row per unique value per column.

    This file is exactly as sensitive as the ORIGINAL, unredacted
    workbook -- anyone holding both it and the redacted file can fully
    de-redact it. It must never be sent alongside the redacted file to
    whatever party the redaction exists to protect against. That warning
    is written into the first few rows of the file itself, since the CSV
    is what actually travels with the case file.

    `carried_from`: the earlier key file(s) loaded for consistent numbering
    ("Use existing key file(s)"). Every row of theirs that this redaction
    didn't already produce is copied in too -- including columns not
    redacted this time -- so this key ALONE un-redacts this workbook and
    every document those keys covered (see _key_lineage.py). Their
    coverage is listed in this key's header.
    """
    rows = list(key_rows)
    if carried_from:
        have = {(r["column"], r["placeholder"]) for r in rows}
        for path in carried_from:
            if not _key_lineage.is_spreadsheet_key(path):
                continue
            for row in _key_lineage.table_rows(path):
                if len(row) < 3 or (row[0], row[1]) in have:
                    continue
                have.add((row[0], row[1]))
                rows.append({"column": row[0], "placeholder": row[1], "value": row[2]})
    with open(key_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["CONFIDENTIAL -- Spreadsheet redaction key"])
        writer.writerow(["Same sensitivity as the ORIGINAL, unredacted workbook -- do NOT send or"])
        writer.writerow(["store this file alongside the redacted workbook. Keep it with your own"])
        writer.writerow(["case file."])
        writer.writerow([f"Source workbook: {os.path.basename(input_path)}"])
        writer.writerow([f"Redacted workbook: {os.path.basename(output_path)}"])
        writer.writerow([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
        _key_lineage.write_coverage_rows(writer, [os.path.basename(input_path)], carried_from)
        writer.writerow([])
        _write_key_body(writer, rows)


def write_key_batch(key_path: str, file_pairs: "list", key_rows: List[dict]) -> None:
    """Same as write_key(), but for a BATCH of workbooks redacted together
    with shared placeholder numbering (see apply_column_redaction()'s
    `state` parameter and base.py's "BATCH REDACTION" section) -- one
    combined key file covering every file in the batch, instead of one
    key file per file.

    `file_pairs` is [(source_basename, redacted_basename), ...] for every
    file in the batch, in the order they were processed; `key_rows` is the
    caller's full accumulated set across every file's apply_column_redaction()
    call (deduplicated by the caller so each (column, value) pair appears
    once, with the placeholder text it was first assigned).

    Same confidentiality as write_key(): this file is exactly as sensitive
    as the ORIGINAL, unredacted workbooks -- anyone holding it plus the
    redacted files can fully de-redact the whole batch.
    """
    with open(key_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["CONFIDENTIAL -- Spreadsheet redaction key (batch)"])
        writer.writerow(["Same sensitivity as the ORIGINAL, unredacted workbooks -- do NOT send or"])
        writer.writerow(["store this file alongside the redacted workbooks. Keep it with your own"])
        writer.writerow(["case file."])
        writer.writerow([f"Files in this batch: {len(file_pairs)}"])
        for source_name, redacted_name in file_pairs:
            writer.writerow([f"  Source: {source_name}  ->  Redacted: {redacted_name}"])
        writer.writerow([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
        _key_lineage.write_coverage_rows(writer, [src for src, _red in file_pairs])
        writer.writerow([])
        _write_key_body(writer, key_rows)


def existing_numbering_from_key(key_path: str) -> dict:
    """Reads a spreadsheet redaction key CSV saved by an EARLIER, SEPARATE
    redaction run (write_key()'s own output, or a combined batch key from
    write_key_batch()) and returns a state mapping usable by
    apply_column_redaction()'s `state` parameter --
    {column_name: {real_value: placeholder_text}}. This is what lets the
    Redact tab's "Use existing key file(s) for consistent numbering"
    picker carry a value's placeholder forward from an earlier,
    separately-redacted workbook into this one, without needing both
    workbooks in the same Batch Redact run.

    Unlike the PDF side, the placeholder TEXT itself (not just a number)
    is carried forward verbatim -- apply_column_redaction() only ever
    assigns a NEW placeholder when a value isn't already present in its
    seeded state, so a value that recurs keeps the exact placeholder
    (including whatever label was chosen when the key file was written);
    only a genuinely new value in this document gets a fresh placeholder
    built from the label chosen THIS time. Use the same placeholder label
    each time you redact a given column if you want every document to
    show identical label text for it.

    Also worth knowing: apply_column_redaction()'s key_rows already
    include every value carried forward via `state`, even ones not
    present in the CURRENT file (see its own docstring) -- so the key CSV
    saved after redacting with an earlier key loaded is itself a superset
    covering both documents, and can be handed forward again next time
    instead of needing to keep collecting every individual key file (the
    PDF side does not have this shortcut -- see
    dentrix_audit_trail.existing_numbering_from_key()'s docstring).

    Raises ValueError if this doesn't look like a spreadsheet key CSV at
    all -- the same "Column, Placeholder, Real Value" header check this
    module's own write_key() output always has. This function is
    intentionally self-contained (a small, local CSV parse) rather than
    importing spreadsheet_unredact_engine.py, matching the sibling
    engine's own preference for not depending across that boundary.
    """
    with open(key_path, "r", newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))

    header_seen = False
    table = []
    for row in rows:
        if not row:
            continue
        if not header_seen:
            if len(row) >= 3 and [c.strip() for c in row[:3]] == ["Column", "Placeholder", "Real Value"]:
                header_seen = True
            continue
        if len(row) >= 3:
            table.append((row[0], row[1], row[2]))

    # A combination's per-part rows ("First Name", <Name 3>, "John") are
    # only there for un-redaction -- numbering is carried by its summary
    # row. Left in, they would map "John" to whichever combination was
    # listed last, and make two genuinely compatible key files look like
    # they conflict.
    combined_part_rows = set()
    for column, placeholder, value in table:
        if is_combined_column_name(column):
            for part in column[len(COMBINED_COLUMN_PREFIX):].split(_COMBINED_PART_SEPARATOR):
                combined_part_rows.add((part, placeholder))

    state = {}
    for column, placeholder, value in table:
        if (column, placeholder) in combined_part_rows:
            continue
        state.setdefault(column, {})[value] = placeholder

    if not header_seen:
        raise ValueError(
            "This doesn't look like a spreadsheet redaction key CSV -- no 'Column, "
            "Placeholder, Real Value' table was found in it. Make sure this is a "
            "*_REDACTION_KEY.csv file saved by an earlier redaction, not some other file "
            "(a PDF name key, for instance, has a different format)."
        )
    return state


def merge_existing_numbering(mappings: List[dict]) -> "tuple[dict, list]":
    """Combines several existing_numbering_from_key() results (one per key
    file the user picked in the Redact tab) into one state mapping, the
    same idea as spreadsheet_unredact_engine.merge_key_mappings()
    combining key files for un-redaction. Returns (merged, conflicts);
    conflicts is [{'column': str, 'value': str, 'placeholders': [p1, p2, ...]}]
    for any (column, value) pair assigned a DIFFERENT placeholder across
    the files given -- this should never happen for key files that
    genuinely belong together, so the caller (main_gui.py) treats any
    conflict as blocking rather than silently picking one placeholder
    over another.
    """
    placeholders_by_key = {}
    for mapping in mappings:
        for column, value_map in mapping.items():
            for value, placeholder in value_map.items():
                seen = placeholders_by_key.setdefault((column, value), [])
                if placeholder not in seen:
                    seen.append(placeholder)
    merged = {}
    conflicts = []
    for (column, value), placeholders in placeholders_by_key.items():
        merged.setdefault(column, {})[value] = placeholders[0]
        if len(placeholders) > 1:
            conflicts.append({"column": column, "value": value, "placeholders": placeholders})
    return merged, conflicts
